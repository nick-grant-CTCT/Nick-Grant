#!/usr/bin/env python3
"""Build the companion stacked-column slide (10 x 5.625 in, CTCT tokens)."""

from __future__ import annotations

import csv
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
CSV_PATH = ROOT / "sales-forward-cumulative.csv"
OUT_PPTX = ROOT / "extra-cumulative-sales.pptx"
CUBE_PATH = ROOT / "ctct-cube.png"

ORANGE = RGBColor(0xFF, 0x8B, 0x00)
GRAY = RGBColor(0x8A, 0x91, 0x96)
DARK = RGBColor(0x59, 0x59, 0x59)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
GRID = RGBColor(0xE6, 0xE6, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CALLOUT_FILL = RGBColor(0xF2, 0xF2, 0xF2)

FONT = "Inter"
Y_MAX = 7000


def load_rows():
    with CSV_PATH.open() as f:
        return list(csv.DictReader(f))


def set_run(run, text, size_pt, bold=False, color=DARK, font=FONT):
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(slide, l, t, w, h, text, size_pt, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size_pt, bold=bold, color=color)
    anchor_map = {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}
    tf._txBody.bodyPr.set("anchor", anchor_map.get(anchor, "t"))
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def line_shape(slide, l, t, w, h, color, weight_pt=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill_shape(shape, color)
    return shape


def disable_shadow(shape):
    spPr = shape._element.spPr
    effect = spPr.find(qn("a:effectLst"))
    if effect is not None:
        spPr.remove(effect)


def main():
    rows = load_rows()
    years = [int(r["Year"]) for r in rows]
    current = [int(r["Current capacity (cumulative)"]) for r in rows]
    extra = [int(r["Extra from added capacity"]) for r in rows]
    added = [c + e for c, e in zip(current, extra)]

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # canvas
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    fill_shape(bg, WHITE)

    add_textbox(
        slide,
        Inches(0.45),
        Inches(0.28),
        Inches(9.1),
        Inches(0.42),
        "Added capacity accumulates extra system sales",
        24,
        bold=True,
        color=DARK,
    )
    add_textbox(
        slide,
        Inches(0.45),
        Inches(0.70),
        Inches(9.1),
        Inches(0.32),
        "Cumulative extra vs the current-capacity path (same rough estimates as the previous slide).",
        12,
        color=MUTED,
    )

    # plot area
    plot_l = Inches(1.15)
    plot_t = Inches(1.20)
    plot_w = Inches(6.55)
    plot_h = Inches(3.55)
    plot_b = plot_t + plot_h

    # y grid + labels
    for val in (0, 2000, 4000, 6000):
        y = plot_b - int(plot_h * val / Y_MAX)
        if val:
            line_shape(slide, plot_l, y, plot_w, Emu(12700), GRID)  # 1px-ish
        add_textbox(
            slide,
            Inches(0.28),
            y - Inches(0.12),
            Inches(0.82),
            Inches(0.24),
            f"{val:,}" if val else "0",
            10,
            color=MUTED,
            align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE,
        )

    add_textbox(
        slide,
        Inches(0.22),
        Inches(1.05),
        Inches(1.4),
        Inches(0.22),
        "System sales",
        10,
        bold=True,
        color=DARK,
    )

    # axes
    line_shape(slide, plot_l, plot_b, plot_w, Emu(19050), DARK)
    line_shape(slide, plot_l, plot_t, Emu(19050), plot_h, DARK)

    n = len(years)
    gap = plot_w / (n * 2.15)
    bar_w = (plot_w - gap * (n + 1)) / n
    # simpler even spacing
    bar_w = Inches(0.95)
    slot = plot_w / n

    for i, (yr, cur, ext, tot) in enumerate(zip(years, current, extra, added)):
        cx = plot_l + int(slot * i + slot / 2)
        left = cx - bar_w // 2
        h_cur = int(plot_h * cur / Y_MAX)
        h_ext = int(plot_h * ext / Y_MAX)
        gray_t = plot_b - h_cur
        org_t = gray_t - h_ext

        gray_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, gray_t, bar_w, h_cur)
        fill_shape(gray_bar, GRAY)
        disable_shadow(gray_bar)

        org_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, org_t, bar_w, max(h_ext, Emu(12700)))
        fill_shape(org_bar, ORANGE)
        disable_shadow(org_bar)

        # total on top of the stack (added-capacity path)
        add_textbox(
            slide,
            left - Inches(0.12),
            org_t - Inches(0.28),
            bar_w + Inches(0.24),
            Inches(0.26),
            f"{tot:,}",
            10,
            color=DARK,
            align=PP_ALIGN.CENTER,
        )
        if ext >= 900:
            add_textbox(
                slide,
                left - Inches(0.08),
                org_t + h_ext // 2 - Inches(0.14),
                bar_w + Inches(0.16),
                Inches(0.28),
                f"+{ext:,}",
                11,
                bold=True,
                color=WHITE,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        else:
            add_textbox(
                slide,
                left + bar_w + Inches(0.04),
                org_t - Inches(0.04),
                Inches(0.7),
                Inches(0.24),
                f"+{ext:,}",
                10,
                bold=True,
                color=ORANGE,
                align=PP_ALIGN.LEFT,
            )

        add_textbox(
            slide,
            left - Inches(0.15),
            plot_b + Inches(0.06),
            bar_w + Inches(0.3),
            Inches(0.24),
            str(yr),
            12,
            color=DARK,
            align=PP_ALIGN.CENTER,
        )

        if yr == 2028:
            add_textbox(
                slide,
                left - Inches(0.25),
                plot_b + Inches(0.28),
                bar_w + Inches(0.5),
                Inches(0.22),
                "MVP (added)",
                9,
                color=ORANGE,
                align=PP_ALIGN.CENTER,
            )
        if yr == 2030:
            add_textbox(
                slide,
                left - Inches(0.35),
                plot_b + Inches(0.28),
                bar_w + Inches(0.7),
                Inches(0.22),
                "Table stakes (added)",
                9,
                color=ORANGE,
                align=PP_ALIGN.CENTER,
            )

    # legend
    legend_l = Inches(7.85)
    legend_t = Inches(1.25)
    sw = Inches(0.22)
    sh = Inches(0.16)
    swatch = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, legend_l, legend_t, sw, sh)
    fill_shape(swatch, GRAY)
    add_textbox(slide, legend_l + Inches(0.28), legend_t - Inches(0.04), Inches(1.7), Inches(0.24), "Current capacity", 11, color=DARK)
    swatch2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, legend_l, legend_t + Inches(0.32), sw, sh)
    fill_shape(swatch2, ORANGE)
    add_textbox(slide, legend_l + Inches(0.28), legend_t + Inches(0.28), Inches(1.7), Inches(0.36), "Extra from added capacity", 11, color=DARK)

    # end-state callout
    call = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(7.75),
        Inches(2.15),
        Inches(1.85),
        Inches(1.15),
    )
    fill_shape(call, CALLOUT_FILL)
    disable_shadow(call)
    tf = call.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    set_run(p.add_run(), "1,000", 22, bold=True, color=ORANGE)
    if len(p.runs) == 2 and not p.runs[0].text:
        p._p.remove(p.runs[0]._r)
    p2 = tf.add_paragraph()
    set_run(p2.add_run(), "extra systems by 2034 vs the current-capacity path", 11, color=DARK)

    notes = (
        "Extra is the gap between the two lines on the previous slide "
        "(Added capacity minus Current capacity), shown as the orange segment. "
        "The gap stays modest through 2030 and compounds after table stakes. "
        "Numbers are rough estimates digitised from that chart, not a new forecast."
    )
    slide.notes_slide.notes_text_frame.text = notes

    add_textbox(slide, Inches(0.45), Inches(5.32), Inches(0.3), Inches(0.22), "6", 10, color=MUTED)

    if CUBE_PATH.exists():
        slide.shapes.add_picture(str(CUBE_PATH), Inches(9.42), Inches(5.18), Inches(0.32), Inches(0.32))

    add_textbox(
        slide,
        Inches(1.15),
        Inches(5.32),
        Inches(7.8),
        Inches(0.22),
        "Rough estimates. Total bar = added-capacity path; gray = current path; orange = extra.",
        8,
        color=MUTED,
    )

    prs.save(OUT_PPTX)
    print(f"Wrote {OUT_PPTX}")


if __name__ == "__main__":
    main()
